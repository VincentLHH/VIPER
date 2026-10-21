#!/usr/bin/env python3
"""Write CPRA-specific HGVSc values into summary-62S.xlsx.

The original case-level HGVSc set is preserved in a new column. The primary
hgvsc column is replaced with the CPRA-specific VEP result. Strict
CPRA-to-case-set assignments and overlap flags remain in separate audit
columns so that unresolved source-data conflicts are not hidden.
"""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


SUMMARY = Path("data/IRD/summary-62S.xlsx")
MAPPING = Path("audit/ird_cpra_hgvsc_row_mapping_20261020.tsv")
BACKUP = Path("data/IRD/summary-62S.before_hgvsc_row_mapping_20261020.xlsx")

AUDIT_COLUMNS = [
    ("hgvsc_case_set_original", "sample_hgvsc_set"),
    ("hgvsc_mapping_status", "mapping_status"),
    ("unassigned_case_hgvsc", "unassigned_case_hgvsc"),
    ("training_overlap_any", "matches_any_training"),
    ("training_overlap_positive", "matches_positive_training"),
]


def find_column(worksheet, field: str) -> int:
    for column in range(1, worksheet.max_column + 1):
        if worksheet.cell(2, column).value == field:
            return column
    raise ValueError(f"Column {field!r} not found")


def display_hgvsc(normalized: str) -> str:
    """Restore conventional HGVSc casing from the normalized match key."""
    if not normalized:
        return ""
    if not normalized.startswith("C."):
        raise ValueError(f"Unexpected normalized HGVSc: {normalized}")
    value = "c." + normalized[2:]
    return re.sub(
        r"(?i)(delins|del|dup|ins|inv)",
        lambda match: match.group(1).lower(),
        value,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--summary", type=Path, default=SUMMARY)
    parser.add_argument("--mapping", type=Path, default=MAPPING)
    parser.add_argument("--backup", type=Path, default=BACKUP)
    args = parser.parse_args()

    if not args.backup.exists():
        shutil.copy2(args.summary, args.backup)

    mapping = pd.read_csv(args.mapping, sep="\t", dtype=str).fillna("")
    key = ["case_id", "cpra"]
    if mapping.duplicated(key).any():
        raise ValueError("Mapping table contains duplicate case_id/cpra keys")
    mapping = mapping.set_index(key)

    workbook = load_workbook(args.summary)
    worksheet = workbook.active
    case_col = find_column(worksheet, "case_ID")
    cpra_col = find_column(worksheet, "cpra")
    hgvsc_col = find_column(worksheet, "hgvsc")

    existing_headers = {
        str(worksheet.cell(2, column).value): column
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(2, column).value
    }
    if "hgvsc_case_set_original" not in existing_headers:
        first_new = worksheet.max_column + 1
        for offset, (header, _) in enumerate(AUDIT_COLUMNS):
            column = first_new + offset
            worksheet.cell(2, column).value = header
            worksheet.cell(2, column).font = Font(bold=True)
            worksheet.column_dimensions[get_column_letter(column)].width = 28
        worksheet.cell(1, first_new).value = "Mapping_audit"
        worksheet.merge_cells(
            start_row=1,
            start_column=first_new,
            end_row=1,
            end_column=first_new + len(AUDIT_COLUMNS) - 1,
        )
        worksheet.cell(1, first_new).alignment = Alignment(horizontal="center")
        worksheet.cell(1, first_new).font = Font(bold=True)
        audit_headers = {
            header: first_new + offset
            for offset, (header, _) in enumerate(AUDIT_COLUMNS)
        }
    else:
        audit_headers = {
            header: existing_headers[header] for header, _ in AUDIT_COLUMNS
        }

    worksheet.cell(2, hgvsc_col).comment = Comment(
        "Row-specific HGVSc selected for the listed CPRA. "
        "The original case-level set is preserved in hgvsc_case_set_original.",
        "VIPER audit",
    )

    updated = 0
    for row in range(3, worksheet.max_row + 1):
        case_id = worksheet.cell(row, case_col).value
        cpra = worksheet.cell(row, cpra_col).value
        if case_id is None or cpra is None:
            continue
        lookup_key = (str(int(case_id)), str(cpra).strip())
        if lookup_key not in mapping.index:
            raise ValueError(f"No mapping for case {lookup_key[0]} / {lookup_key[1]}")

        record = mapping.loc[lookup_key]
        cpra_hgvsc = str(record["cpra_hgvsc"])
        if not cpra_hgvsc:
            raise ValueError(f"No CPRA-specific HGVSc for {lookup_key}")

        case_set_col = audit_headers["hgvsc_case_set_original"]
        if not worksheet.cell(row, case_set_col).value:
            worksheet.cell(row, case_set_col).value = worksheet.cell(
                row, hgvsc_col
            ).value

        worksheet.cell(row, hgvsc_col).value = display_hgvsc(cpra_hgvsc)
        for header, source_field in AUDIT_COLUMNS[1:]:
            value = record[source_field]
            if source_field == "unassigned_case_hgvsc" and value:
                value = ";".join(
                    display_hgvsc(token)
                    for token in str(value).split(";")
                    if token
                )
            if source_field in {
                "matches_any_training",
                "matches_positive_training",
            }:
                value = str(value).strip().lower() == "true"
            worksheet.cell(row, audit_headers[header]).value = value
        updated += 1

    workbook.save(args.summary)
    print(f"Updated {updated} IRD rows in {args.summary}")
    print(f"Backup: {args.backup}")


if __name__ == "__main__":
    main()
