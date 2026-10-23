import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from scripts.bootstrap_topk_ci import (
    build_patient_bootstrap_indices,
    compute_cohort_tables,
    compute_iei_overlap_mask,
    filter_non_overlapping,
    load_iei_s2,
    load_ird_summary,
    paired_variant_topk_difference,
    patient_topk_ci,
    variant_topk_ci,
    write_cohort_tables,
)


def _rank_data():
    return pd.DataFrame(
        {
            "patient_id": ["P1", "P1", "P2", "P2", "P3"],
            "VIPER": [1, 20, 2, 4, np.nan],
            "GeneT": [25, 3, 30, 2, np.nan],
        }
    )


def test_bootstrap_indices_are_deterministic_and_patient_level():
    patient_ids = [f"P{i}" for i in range(1, 21)]
    first = build_patient_bootstrap_indices(patient_ids, n_resamples=100, seed=42)
    second = build_patient_bootstrap_indices(patient_ids, n_resamples=100, seed=42)

    assert first.shape == (100, 20)
    np.testing.assert_array_equal(first, second)
    assert first.min() == 0
    assert first.max() == 19


def test_variant_topk_ci_counts_missing_ranks_as_failures():
    data = _rank_data()
    indices = build_patient_bootstrap_indices(
        sorted(data["patient_id"].unique()), n_resamples=200, seed=42
    )

    result = variant_topk_ci(data, rank_column="VIPER", k=10, indices=indices)

    assert result.estimate == pytest.approx(0.6)
    assert result.low <= result.estimate <= result.high


def test_patient_level_any_and_complete_are_distinct():
    data = _rank_data()
    indices = build_patient_bootstrap_indices(
        sorted(data["patient_id"].unique()), n_resamples=200, seed=42
    )

    any_result = patient_topk_ci(
        data, rank_column="VIPER", k=10, indices=indices, mode="any"
    )
    complete_result = patient_topk_ci(
        data, rank_column="VIPER", k=10, indices=indices, mode="complete"
    )

    assert any_result.estimate == pytest.approx(2 / 3)
    assert complete_result.estimate == pytest.approx(1 / 3)


def test_paired_difference_uses_identical_patient_draws():
    data = _rank_data()
    indices = build_patient_bootstrap_indices(
        sorted(data["patient_id"].unique()), n_resamples=200, seed=42
    )

    result = paired_variant_topk_difference(
        data,
        reference_column="VIPER",
        comparator_column="GeneT",
        k=10,
        indices=indices,
    )

    assert result.estimate == pytest.approx(0.2)
    assert result.low <= result.estimate <= result.high


def _write_grouped_sheet(path, sheet_name, group_row, field_row, data_rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(group_row)
    worksheet.append(field_row)
    for row in data_rows:
        worksheet.append(row)
    workbook.save(path)


def _write_grouped_sheet_with_title_and_note(
    path, sheet_name, group_row, field_row, data_rows
):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append([f"Table {sheet_name}. Evaluation data."])
    worksheet.append(["This row is a table note and must not be parsed as data."])
    worksheet.append(group_row)
    worksheet.append(field_row)
    for row in data_rows:
        worksheet.append(row)
    workbook.save(path)


def test_load_iei_s2_maps_model_blocks_and_missing_ranks(tmp_path):
    path = tmp_path / "iei.xlsx"
    _write_grouped_sheet(
        path,
        "S2",
        [
            "Sample_info",
            None,
            None,
            "VIPER",
            None,
            None,
            "Xrare",
            None,
            None,
            None,
        ],
        [
            "case_ID",
            "cpra",
            "Gene",
            "label",
            "score",
            "rank",
            "label",
            "score",
            "rank_obs",
            "rank_aug",
        ],
        [
            ["P1", "chr1-1-A-C", "GENE1", "positive", 0.9, 2, "—", 0, "#N/A", 8],
        ],
    )

    data = load_iei_s2(path)

    assert list(data.columns) == [
        "patient_id",
        "cpra",
        "gene",
        "VIPER",
        "VIPER_positive",
        "Xrare observed",
        "Xrare rank-augmented",
    ]
    assert data.loc[0, "patient_id"] == "P1"
    assert data.loc[0, "VIPER"] == 2
    assert np.isnan(data.loc[0, "Xrare observed"])
    assert data.loc[0, "Xrare rank-augmented"] == 8


def test_load_iei_s2_tolerates_title_and_note_rows(tmp_path):
    path = tmp_path / "iei_with_title.xlsx"
    _write_grouped_sheet_with_title_and_note(
        path,
        "S2",
        [
            "Sample_info",
            None,
            None,
            "VIPER",
            None,
            None,
            "GeneT",
            None,
            None,
        ],
        [
            "case_ID",
            "cpra",
            "Gene",
            "label",
            "score",
            "rank",
            "label",
            "score",
            "rank",
        ],
        [
            ["P1", "chr1-1-A-C", "GENE1", "positive", 0.9, 2, "positive", 0.8, 3],
        ],
    )

    data = load_iei_s2(path)

    assert data.to_dict("records") == [
        {
            "patient_id": "P1",
            "cpra": "chr1-1-A-C",
            "gene": "GENE1",
            "VIPER": 2,
            "VIPER_positive": True,
            "GeneT": 3,
        }
    ]


def test_load_grouped_sheet_rejects_missing_required_header(tmp_path):
    path = tmp_path / "invalid.xlsx"
    _write_grouped_sheet(
        path,
        "S2",
        ["Sample_info", "VIPER"],
        ["patient_id", "rank"],
        [["P1", 1]],
    )

    with pytest.raises(KeyError, match="case_ID"):
        load_iei_s2(path)


def test_load_ird_summary_maps_observed_rank_and_overlap_flag(tmp_path):
    path = tmp_path / "ird.xlsx"
    _write_grouped_sheet(
        path,
        "Sheet1",
        [
            "Sample_Info",
            None,
            "VIPER",
            None,
            None,
            "Xrare",
            None,
            None,
            None,
            "Mapping_audit",
        ],
        [
            "case_ID",
            "cpra",
            "label",
            "score",
            "rank",
            "label",
            "score",
            "rank_ori",
            "rank_aug",
            "training_overlap_any",
        ],
        [
            ["P1", "chr1-1-A-C", "positive", 0.9, 2, "—", 0, 4, 3, True],
        ],
    )

    data = load_ird_summary(path)

    assert list(data.columns) == [
        "patient_id",
        "cpra",
        "VIPER",
        "VIPER_positive",
        "Xrare observed",
        "Xrare rank-augmented",
        "training_overlap_any",
    ]
    assert data.loc[0, "Xrare observed"] == 4
    assert bool(data.loc[0, "training_overlap_any"])


def test_compute_cohort_tables_preserves_metric_hierarchy():
    data = _rank_data()
    data["VIPER_positive"] = [True, False, True, True, False]

    tables = compute_cohort_tables(data, n_resamples=200, seed=42)

    summary = tables["model_summary"].set_index(["model", "k"])
    paired = tables["paired_top10"].set_index("comparison")
    complete = tables["patient_complete_top10"].set_index("model")

    assert summary.loc[("VIPER", 10), "estimate"] == pytest.approx(0.6)
    assert paired.loc["VIPER minus GeneT", "estimate"] == pytest.approx(0.2)
    assert complete.loc["VIPER", "estimate"] == pytest.approx(1 / 3)


def test_filter_non_overlapping_uses_explicit_training_flag():
    data = pd.DataFrame(
        {
            "patient_id": ["P1", "P2"],
            "VIPER": [1, 2],
            "training_overlap_any": [True, False],
        }
    )

    filtered = filter_non_overlapping(data)

    assert filtered["patient_id"].tolist() == ["P2"]


def test_compute_iei_overlap_mask_uses_cpra_and_gene_hgvsc(tmp_path):
    training = tmp_path / "training.jsonl"
    prompt = (
        "The basic information of variant chr1-10-A-C is as follows:\n"
        "Gene symbol: GENE2\nTranscript: NM_2.1\nDNA change: c.2A>C\n"
        "Protein change: p.Lys1Thr"
    )
    training.write_text(
        json.dumps({"messages": [{"role": "assistant", "content": prompt}]}) + "\n",
        encoding="utf-8",
    )

    metadata_path = tmp_path / "metadata.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["cpra", "gene symbol", "dna change", "protein change"])
    worksheet.append(["chr1-10-A-C", "GENE1", "c.10A>C", "p.Lys4Thr"])
    worksheet.append(["chr1-20-A-G", "GENE2", "c.2A>C", "p.Lys1Thr"])
    worksheet.append(["chr1-30-A-T", "GENE3", "c.30A>T", "p.Lys10Phe"])
    workbook.save(metadata_path)

    data = pd.DataFrame(
        {
            "patient_id": ["P1", "P2", "P3"],
            "cpra": ["chr1-10-A-C", "chr1-20-A-G", "chr1-30-A-T"],
            "gene": ["GENE1", "GENE2", "GENE3"],
        }
    )

    mask = compute_iei_overlap_mask(
        data,
        training_jsonl=training,
        metadata_xlsx=metadata_path,
    )

    assert mask.tolist() == [True, True, False]


def test_write_cohort_tables_creates_machine_readable_outputs(tmp_path):
    data = _rank_data()
    data["VIPER_positive"] = [True, False, True, True, False]
    tables = compute_cohort_tables(data, n_resamples=200, seed=42)

    written = write_cohort_tables(tables, output_dir=tmp_path, prefix="S3")

    assert set(written) == {
        "model_summary",
        "paired_topk",
        "patient_complete_top10",
        "label_metrics",
    }
    assert (tmp_path / "S3a_model_ci.tsv").exists()
    assert (tmp_path / "S3b_paired_ci.tsv").exists()
    assert (tmp_path / "S3c_patient_complete_ci.tsv").exists()
    assert (tmp_path / "S3_label_metrics_ci.tsv").exists()
